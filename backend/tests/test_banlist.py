from collections import Counter
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
import pytest
from fastapi.testclient import TestClient

from app import models  # noqa: F401
from app.db.bootstrap import seed_defaults
from app.db.session import Base, SessionLocal, engine
from app.models import BanImport, Event, Song, User
from app.modules.banlist.service import parse_ban_workbook, similarity
from sqlalchemy import select
from app.main import app


@pytest.fixture(autouse=True)
def reset_db():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    with SessionLocal() as db:
        seed_defaults(db)


@pytest.fixture
def client():
    with TestClient(app) as test_client:
        yield test_client


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(1, 1, "#1 - Pure")
    sheet.cell(2, 1, "曲名")
    sheet.cell(2, 2, "作者")
    sheet.cell(2, 3, "备注")
    sheet.cell(3, 1, "Ban Song")
    sheet.cell(3, 2, "Artist")
    sheet.cell(3, 3, "历史备注")
    sheet.cell(1, 5, "#2 - Pure Plus")
    sheet.cell(2, 5, "曲名")
    sheet.cell(2, 6, "作者")
    sheet.cell(2, 7, "备注")
    sheet.cell(3, 5, "Second Song")
    sheet.cell(3, 6, "Second Artist")
    sheet.cell(4, 5, "Chaotic Ørder")
    sheet.cell(4, 6, "TAG VS Kai")
    sheet.cell(1, 9, "#3 - Devour")
    sheet.cell(2, 9, "曲名")
    sheet.cell(2, 10, "作者")
    sheet.cell(2, 11, "备注")
    sheet.cell(3, 9, "Third Song")
    sheet.cell(3, 10, "Third Artist")
    sheet.cell(1, 13, "#4 - Devour plus")
    sheet.cell(2, 13, "曲名")
    sheet.cell(2, 14, "作者")
    sheet.cell(2, 15, "备注")
    sheet.cell(3, 13, "Fourth Song")
    sheet.cell(3, 14, "Fourth Artist")
    sheet.cell(3, 15, "Fourth Note")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_similarity_uses_rapidfuzz_wratio_and_rounds() -> None:
    with patch("app.modules.banlist.service.fuzz.WRatio", return_value=87.654) as wratio:
        assert similarity("normalized-title", "normalized-artist") == 87.7
        wratio.assert_called_once_with("normalized-title", "normalized-artist")

    with patch("app.modules.banlist.service.fuzz.WRatio") as wratio:
        assert similarity("", "non-empty") == 0.0
        wratio.assert_not_called()


def test_parser_discovers_future_blocks_and_preserves_validation_issues() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(1, 17, "#5 - Future Round")
    sheet.cell(2, 17, "song_name")
    sheet.cell(2, 18, "artist")
    sheet.cell(2, 19, "note")
    sheet.cell(3, 17, "Future Song")
    sheet.cell(3, 18, "Future Artist")
    sheet.cell(3, 19, "Future Note")
    sheet.cell(4, 17, "Future Song")
    sheet.cell(4, 18, "Future Artist")
    sheet.cell(5, 17, "Missing Artist")
    output = BytesIO()
    workbook.save(output)

    entries, issues = parse_ban_workbook(output.getvalue())

    assert [(entry.round_label, entry.song_name, entry.artist, entry.remark) for entry in entries] == [
        ("#5 - Future Round", "Future Song", "Future Artist", "Future Note"),
        ("#5 - Future Round", "Future Song", "Future Artist", ""),
    ]
    assert len(issues) == 2
    assert any("缺少曲名或作者" in issue for issue in issues)
    assert any("与同届已有曲目重复" in issue for issue in issues)


def test_parser_rejects_empty_and_unrecognized_workbooks() -> None:
    with pytest.raises(ValueError, match="文件为空"):
        parse_ban_workbook(b"")

    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(1, 1, "#1 - Invalid")
    sheet.cell(2, 1, "不是曲名表头")
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(ValueError, match="未找到有效的 Ban 曲区块"):
        parse_ban_workbook(output.getvalue())


def test_bundled_number_four_workbook_parses_completely_and_is_downloadable(client) -> None:
    workbook_path = next((Path(__file__).resolve().parents[2] / "banlist").glob("*#4*.xlsx"))
    entries, issues = parse_ban_workbook(workbook_path.read_bytes())

    assert issues == []
    assert Counter(entry.round_label for entry in entries) == {
        "#1 - Pure": 21,
        "#2 - Pure Plus": 50,
        "#3 - Devour": 47,
        "#4 - Devour plus": 50,
    }
    assert len(entries) == 168

    metadata = client.get("/api/v1/assets/banlist")
    assert metadata.status_code == 200
    assert metadata.json()["file_name"] == workbook_path.name
    download = client.get("/api/v1/assets/banlist/download")
    assert download.status_code == 200
    assert download.content == workbook_path.read_bytes()


def register(client, user_code: str, identity: str = "participant") -> None:
    response = client.post(
        "/api/v1/auth/register",
        json={"user_code": user_code, "qq_id": user_code, "password": "secret123", "identity": identity},
    )
    assert response.status_code == 201, response.text


def login(client, user_code: str, password: str = "secret123") -> None:
    response = client.post("/api/v1/auth/login", json={"user_code": user_code, "password": password})
    assert response.status_code == 200, response.text


def login_admin(client) -> None:
    login(client, "admin", "change-me-please")


def test_ban_import_preview_publish_and_matching(client):
    login_admin(client)
    uploaded = client.post(
        "/api/v1/admin/banlist/import",
        files={"file": ("ban.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert uploaded.status_code == 200, uploaded.text
    preview = uploaded.json()
    assert preview["entry_count"] == 5
    assert preview["issues"] == []
    assert {entry["round"] for entry in preview["entries"]} == {
        "#1 - Pure",
        "#2 - Pure Plus",
        "#3 - Devour",
        "#4 - Devour plus",
    }

    published = client.post(f"/api/v1/admin/banlist/{preview['id']}/publish")
    assert published.status_code == 200, published.text
    assert published.json()["status"] == "published"

    register(client, "checker")
    login(client, "checker")
    exact = client.post("/api/v1/banlist/check", json={"title": "Ban Song", "artist": "Artist"})
    assert exact.status_code == 200
    assert exact.json()["status"] == "exact"
    assert exact.json()["matches"][0]["round"] == "#1 - Pure"

    fourth = client.post("/api/v1/banlist/check", json={"title": "Fourth Song", "artist": "Fourth Artist"})
    assert fourth.status_code == 200
    assert fourth.json()["status"] == "exact"
    assert fourth.json()["matches"][0]["round"] == "#4 - Devour plus"

    fourth_search = client.get("/api/v1/banlist/search", params={"title": "Fourth Song"})
    assert fourth_search.status_code == 200
    assert fourth_search.json()["items"][0]["title"] == "Fourth Song"
    assert fourth_search.json()["items"][0]["round"] == "#4 - Devour plus"

    normalized = client.post("/api/v1/banlist/check", json={"title": "Ban-Song", "artist": "artist"})
    assert normalized.status_code == 200
    assert normalized.json()["status"] == "exact"

    review = client.post("/api/v1/banlist/check", json={"title": "Ban Sng", "artist": "Artist"})
    assert review.status_code == 200
    assert review.json()["status"] == "review"

    # The manual search scores this pair at 95%, so the automatic check must
    # not report the same complete input as clear just because the title has
    # one stylized character.
    stylized_title = client.post(
        "/api/v1/banlist/check",
        json={"title": "chaotic order", "artist": "tag vs kai"},
    )
    assert stylized_title.status_code == 200
    assert stylized_title.json()["status"] == "review"
    assert stylized_title.json()["matches"][0]["score"] == 95.0

    stylized_search = client.get(
        "/api/v1/banlist/search",
        params={"title": "chaotic order", "artist": "tag vs kai"},
    )
    assert stylized_search.status_code == 200
    assert stylized_search.json()["items"][0]["score"] == 95.0

    lower_threshold = client.post(
        "/api/v1/banlist/check",
        json={"title": "chaotic ordr", "artist": "tag vs kai"},
    )
    assert lower_threshold.status_code == 200
    assert lower_threshold.json()["status"] == "review"
    assert lower_threshold.json()["matches"][0]["score"] == 92.2

    search = client.get("/api/v1/banlist/search", params={"title": "Second"})
    assert search.status_code == 200
    assert search.json()["items"][0]["title"] == "Second Song"


def test_ban_check_requires_login_and_song_pool_rechecks(client):
    unauthenticated = client.post("/api/v1/banlist/check", json={"title": "Ban Song", "artist": "Artist"})
    assert unauthenticated.status_code == 401

    login_admin(client)
    uploaded = client.post(
        "/api/v1/admin/banlist/import",
        files={"file": ("ban.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    import_id = uploaded.json()["id"]
    assert client.post(f"/api/v1/admin/banlist/{import_id}/publish").status_code == 200

    register(client, "pool-user")
    login(client, "pool-user")
    blocked = client.post(
        "/api/v1/song-pool/me",
        json={"song_name": "Ban Song", "artist": "Artist", "song_type": "A", "remark": ""},
    )
    assert blocked.status_code == 409
    assert "命中往届 Ban 曲" in blocked.json()["detail"]

    review_blocked = client.post(
        "/api/v1/song-pool/me",
        json={"song_name": "Ban Sng", "artist": "Artist", "song_type": "A", "remark": ""},
    )
    assert review_blocked.status_code == 409

    allowed = client.post(
        "/api/v1/song-pool/me",
        json={"song_name": "Ban Sng", "artist": "Artist", "song_type": "A", "remark": "", "acknowledge_ban_warning": True},
    )
    assert allowed.status_code == 200, allowed.text

    with SessionLocal() as db:
        event = db.scalar(select(Event).where(Event.is_current.is_(True)))
        user = db.scalar(select(User).where(User.user_code == "pool-user"))
        assert event is not None and user is not None
        event.settings.phase_mode = "manual"
        event.settings.manual_phase = "submission_1"
        db.add(Song(event_id=event.id, submitted_by_id=user.id, song_name="Ban Song", artist="Artist", song_type="A", remark=""))
        db.commit()
        song_id = db.scalar(select(Song.id).where(Song.song_name == "Ban Song", Song.submitted_by_id == user.id))

    blocked_submission = client.post(
        "/api/v1/submissions",
        data={"song_id": str(song_id), "track": "normal"},
        files={"file": ("ignored.zip", b"not parsed because Ban is checked first", "application/zip")},
    )
    assert blocked_submission.status_code == 409
    assert "命中往届 Ban 曲" in blocked_submission.json()["detail"]

    with SessionLocal() as db:
        assert db.scalar(select(BanImport).where(BanImport.status == "published")) is not None
